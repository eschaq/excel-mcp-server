"""Build evals/fixtures/eval_workbook.xlsx. evals/prepare.py calls build(); see "Evals" in the README.

Sheets:
  Orders     20 orders: dates, text, integers, decimals, a formula column, and deliberate blanks
  Customers  8 customers: text, dates, booleans, one blank email
  Budget     a small budget with SUM / AVERAGE / IF formulas
  Notes      free text for search tests
  Empty      no cells at all
"""

import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ORDERS = [
    # id, date, customer, region, product, qty, unit price
    (1001, dt.date(2026, 1, 4), "Acme Corp", "West", "Widget", 12, 2.50),
    (1002, dt.date(2026, 1, 9), "Globex", "East", "Gadget", 4, 19.99),
    (1003, dt.date(2026, 1, 15), "Initech", "North", "Gizmo", 30, 7.25),
    (1004, dt.date(2026, 1, 22), "Acme Corp", "West", "Gadget", 8, 19.99),
    (1005, dt.date(2026, 2, 3), "Umbrella", "South", "Widget", 100, 2.50),
    (1006, dt.date(2026, 2, 11), "Globex", "East", "Doohickey", 6, 45.00),
    (1007, dt.date(2026, 2, 18), "Hooli", "West", "Gizmo", None, 7.25),  # missing qty
    (1008, dt.date(2026, 2, 27), "Initech", "North", "Widget", 55, 2.50),
    (1009, dt.date(2026, 3, 5), "Acme Corp", "West", "Doohickey", 3, 45.00),
    (1010, dt.date(2026, 3, 12), "Stark Industries", "East", "Gadget", 25, 19.99),
    (1011, dt.date(2026, 3, 19), "Umbrella", "South", "Gizmo", 14, 7.25),
    (1012, dt.date(2026, 3, 30), "Hooli", "West", "Widget", 40, 2.50),
    (1013, dt.date(2026, 4, 2), "Wayne Enterprises", None, "Doohickey", 10, 45.00),  # missing region
    (1014, dt.date(2026, 4, 9), "Globex", "East", "Widget", 75, 2.50),
    (1015, dt.date(2026, 4, 16), "Stark Industries", "East", "Doohickey", 18, 45.00),
    (1016, dt.date(2026, 4, 23), "Initech", "North", "Gadget", 9, 19.99),
    (1017, dt.date(2026, 5, 1), "Acme Corp", "West", "Gizmo", 22, 7.25),
    (1018, dt.date(2026, 5, 8), "Umbrella", "South", "Gadget", 2, 19.99),
    (1019, dt.date(2026, 5, 15), "Wayne Enterprises", "North", "Widget", 60, 2.50),
    (1020, dt.date(2026, 5, 22), "Hooli", "West", "Doohickey", 5, 45.00),
]

CUSTOMERS = [
    ("Acme Corp", "orders@acme.example", "West", dt.date(2021, 3, 14), True),
    ("Globex", "buying@globex.example", "East", dt.date(2022, 7, 1), True),
    ("Initech", "tps@initech.example", "North", dt.date(2019, 11, 20), False),
    ("Umbrella", "procurement@umbrella.example", "South", dt.date(2023, 1, 9), True),
    ("Hooli", None, "West", dt.date(2024, 5, 30), False),  # missing email
    ("Stark Industries", "supply@stark.example", "East", dt.date(2020, 9, 2), True),
    ("Wayne Enterprises", "purchasing@wayne.example", "North", dt.date(2025, 2, 17), True),
    ("Soylent", "hello@soylent.example", "South", dt.date(2026, 1, 5), False),
]


def bold_header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def build(path: Path) -> None:
    wb = openpyxl.Workbook()

    orders = wb.active
    orders.title = "Orders"
    bold_header(orders, ["Order ID", "Date", "Customer", "Region", "Product", "Qty", "Unit Price", "Total"])
    for r, row in enumerate(ORDERS, start=2):
        orders.append(list(row) + [f"=F{r}*G{r}"])
        orders.cell(row=r, column=2).number_format = "yyyy-mm-dd"

    cust = wb.create_sheet("Customers")
    bold_header(cust, ["Customer", "Email", "Region", "Customer Since", "Active"])
    for row in CUSTOMERS:
        cust.append(list(row))
        cust.cell(row=cust.max_row, column=4).number_format = "yyyy-mm-dd"

    budget = wb.create_sheet("Budget")
    bold_header(budget, ["Category", "Q1", "Q2", "Total", "Over Limit?"])
    for r, (cat, q1, q2) in enumerate(
        [("Marketing", 12000, 15500), ("Payroll", 88000, 91000), ("Software", 4200, 3900), ("Travel", 6500, 8100)],
        start=2,
    ):
        budget.append([cat, q1, q2, f"=SUM(B{r}:C{r})", f'=IF(D{r}>20000,"Yes","No")'])
    budget.append(["Grand Total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)", None])
    budget.append(["Average", "=AVERAGE(B2:B5)", "=AVERAGE(C2:C5)", None, None])

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Q2 planning notes"
    notes["A3"] = "Acme Corp asked for a volume discount on Widgets."
    notes["A4"] = "Globex contract renewal due 2026-06-30."
    notes["A5"] = "Invoice INV-2026-0042 disputed by Hooli."
    notes["C7"] = 1005  # a bare number that also appears as an order ID

    wb.create_sheet("Empty")
    wb.save(path)


if __name__ == "__main__":
    build(Path(__file__).with_name("eval_workbook.xlsx"))
