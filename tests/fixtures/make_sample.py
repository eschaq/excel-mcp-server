"""Regenerate tests/fixtures/sample.xlsx: python tests/fixtures/make_sample.py"""

import datetime as dt
from pathlib import Path

import openpyxl

SALES = [
    ("Region", "Rep", "Product", "Units", "Price", "Date", "Revenue"),
    ("West", "Alice", "Widget", 10, 2.5, dt.date(2026, 1, 5)),
    ("East", "Bob", "Gadget", 4, 10.0, dt.date(2026, 1, 9)),
    ("West", "Carol", "Gizmo", 7, 6.0, dt.date(2026, 2, 1)),
    ("North", "Dave", "Widget", 12, 2.5, dt.date(2026, 2, 14)),
    ("East", "Alice", "Gizmo", 3, 6.0, dt.date(2026, 3, 2)),
    ("South", "Erin", "Gadget", None, 10.0, dt.date(2026, 3, 20)),
]


def build(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(SALES[0])
    for i, row in enumerate(SALES[1:], start=2):
        ws.append(list(row) + [f"=D{i}*E{i}"])
    ws.append([])
    ws.append(["Total", None, None, "=SUM(D2:D7)", None, None, "=SUM(G2:G7)"])

    inv = wb.create_sheet("Inventory")
    inv.append(["SKU", "Product", "In Stock"])
    inv.append(["W-1", "Widget", 120])
    inv.append(["G-2", "Gadget", 0])
    inv.append(["Z-3", "Gizmo", 45])
    wb.save(path)


def build_xls(path: Path) -> None:
    """Legacy .xls copy of the Sales sheet. Needs xlwt (pip install xlwt), a test-only tool."""
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sales")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for r, row in enumerate(SALES):
        for c, v in enumerate(row):
            if v is not None:
                ws.write(r, c, v, date_style) if isinstance(v, dt.date) else ws.write(r, c, v)
    wb.save(str(path))


if __name__ == "__main__":
    build(Path(__file__).with_name("sample.xlsx"))
    build_xls(Path(__file__).with_name("sample.xls"))
