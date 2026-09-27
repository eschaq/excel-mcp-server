import asyncio

import openpyxl
import pytest
from mcp import Client

from excel_mcp.excel_ops import Config, ExcelOpsError
from excel_mcp.tools import build_server

WRITE_TOOLS = {"write_cell", "write_range", "create_sheet", "create_workbook"}
READ_TOOLS = {"read_spreadsheet", "get_sheet_data", "get_formulas", "search", "get_summary_stats", "apply_filter"}


# ---- write_cell / write_range ---------------------------------------------- #

def test_write_cell_value_and_formula(ops, workdir):
    out = ops.write_cell("sample.xlsx", "$b$2", "Alicia")
    assert out["old_value"] == "Alice" and out["cell"] == "B2"
    ops.write_cell("sample.xlsx", "H2", "=G2*2", sheet="Sales")

    ws = openpyxl.load_workbook(workdir / "sample.xlsx")["Sales"]
    assert ws["B2"].value == "Alicia"
    assert ws["H2"].value == "=G2*2"
    assert ws["G3"].value == "=D3*E3"  # untouched formulas survive the round trip


def test_write_cell_clears_with_none(ops, workdir):
    ops.write_cell("sample.xlsx", "A2", None)
    assert openpyxl.load_workbook(workdir / "sample.xlsx")["Sales"]["A2"].value is None


def test_write_range(ops, workdir):
    out = ops.write_range("sample.xlsx", "B6", [["X", "Y"], [1, 2, 3]], sheet="Inventory")
    assert out["range"] == "B6:D7" and out["cells_written"] == 5
    ws = openpyxl.load_workbook(workdir / "sample.xlsx")["Inventory"]
    assert [[c.value for c in row] for row in ws["B6:D7"]] == [["X", "Y", None], [1, 2, 3]]


def test_write_rejects_bad_input(ops):
    with pytest.raises(ExcelOpsError, match="Invalid cell"):
        ops.write_cell("sample.xlsx", "not-a-cell", 1)
    with pytest.raises(ExcelOpsError, match="outside Excel's grid"):
        ops.write_cell("sample.xlsx", "XFE1", 1)
    with pytest.raises(ExcelOpsError, match="Invalid range"):
        ops.get_formulas("sample.xlsx", "A1:??")
    with pytest.raises(ExcelOpsError, match="2D array"):
        ops.write_range("sample.xlsx", "A1", [])
    with pytest.raises(ExcelOpsError, match="Unsupported file type"):
        ops.write_cell("sales.csv", "A1", 1)
    with pytest.raises(ExcelOpsError, match="Unsupported file type"):
        ops.write_cell("sample.xls", "A1", 1)


# ---- create_sheet / create_workbook ---------------------------------------- #

def test_create_sheet(ops, workdir):
    out = ops.create_sheet("sample.xlsx", "Summary", headers=["Metric", "Value"], index=0)
    assert out["sheets"] == ["Summary", "Sales", "Inventory"]
    ws = openpyxl.load_workbook(workdir / "sample.xlsx")["Summary"]
    assert ws["A1"].value == "Metric" and ws["A1"].font.bold

    with pytest.raises(ExcelOpsError, match="already exists"):
        ops.create_sheet("sample.xlsx", "sales")
    with pytest.raises(ExcelOpsError, match="Invalid sheet name"):
        ops.create_sheet("sample.xlsx", "bad/name")


def test_create_workbook(ops, workdir):
    out = ops.create_workbook("reports/q1.xlsx", sheet_name="Q1", headers=["Date", "Amount"])
    assert out["sheets"] == ["Q1"]
    assert ops.read_spreadsheet("reports/q1.xlsx")["sheets"][0]["preview"] == [["Date", "Amount"]]

    with pytest.raises(ExcelOpsError, match="already exists"):
        ops.create_workbook("reports/q1.xlsx")
    ops.create_workbook("reports/q1.xlsx", overwrite=True)
    with pytest.raises(ExcelOpsError, match="only creates .xlsx"):
        ops.create_workbook("macro.xlsm")


def test_create_workbook_outside_sandbox(ops, workdir):
    with pytest.raises(ExcelOpsError, match="Access denied"):
        ops.create_workbook(str(workdir.parent / "escape.xlsx"))


def test_no_temp_files_left_behind(ops, workdir):
    ops.write_cell("sample.xlsx", "A2", "x")
    assert sorted(p.name for p in workdir.iterdir()) == ["sales.csv", "sample.xls", "sample.xlsx"]


# ---- read-only mode --------------------------------------------------------- #

def test_read_only_blocks_writes(ro_ops, workdir):
    before = (workdir / "sample.xlsx").read_bytes()
    with pytest.raises(ExcelOpsError, match="read-only"):
        ro_ops.write_cell("sample.xlsx", "A1", "x")
    with pytest.raises(ExcelOpsError, match="read-only"):
        ro_ops.create_workbook("new.xlsx")
    assert (workdir / "sample.xlsx").read_bytes() == before
    assert ro_ops.read_spreadsheet("sample.xlsx")["sheet_count"] == 2


# ---- through the MCP protocol ---------------------------------------------- #

def _run(coro):
    return asyncio.run(coro)


async def _tool_names(server):
    async with Client(server) as client:
        return {t.name for t in (await client.list_tools()).tools}


def test_mcp_lists_all_tools(workdir):
    assert _run(_tool_names(build_server(Config(allowed_dirs=[workdir])))) == READ_TOOLS | WRITE_TOOLS


def test_mcp_read_only_hides_write_tools(workdir):
    assert _run(_tool_names(build_server(Config(allowed_dirs=[workdir], read_only=True)))) == READ_TOOLS


def test_mcp_call_and_error_message(workdir):
    server = build_server(Config(allowed_dirs=[workdir]))

    async def go():
        async with Client(server) as client:
            ok = await client.call_tool("apply_filter", {
                "path": "sample.xlsx",
                "conditions": [{"column": "Region", "op": "eq", "value": "East"}],
                "columns": ["Rep"],
            })
            bad = await client.call_tool("read_spreadsheet", {"path": "../../etc/passwd.xlsx"})
            return ok, bad

    ok, bad = _run(go())
    assert not ok.is_error and '"Bob"' in ok.content[0].text
    assert bad.is_error and "Access denied" in bad.content[0].text
