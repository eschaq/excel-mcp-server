"""Core spreadsheet operations for the Excel MCP server.

Everything here is plain Python with no MCP dependency, so it can be tested and
reused directly. Every public method on ``ExcelOps`` validates its path against
the configured sandbox, enforces the file size limit, and logs the operation.
"""

from __future__ import annotations

import csv
import datetime as dt
import logging
import math
import os
import re
import statistics
import tempfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Literal

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils.exceptions import CellCoordinatesException
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
    range_boundaries,
)

log = logging.getLogger("excel_mcp")

READ_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".csv"}
WRITE_EXTENSIONS = {".xlsx", ".xlsm"}
MAX_ROWS_PER_CALL = 5000
EXCEL_MAX_ROW, EXCEL_MAX_COL = 1_048_576, 16_384
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

FilterOp = Literal[
    "eq", "ne", "gt", "gte", "lt", "lte", "between",
    "contains", "not_contains", "startswith", "endswith", "regex",
    "in", "not_in", "is_null", "not_null",
]


class ExcelOpsError(Exception):
    """An anticipated failure whose message is safe and useful to show the model."""


@dataclass
class Config:
    allowed_dirs: list[Path] = field(default_factory=lambda: [Path.home()])
    read_only: bool = False
    max_file_mb: float = 50.0

    def __post_init__(self) -> None:
        if not self.allowed_dirs:
            raise ValueError("At least one allowed directory is required")
        self.allowed_dirs = [Path(d).expanduser().resolve() for d in self.allowed_dirs]


@dataclass
class Condition:
    column: str
    op: FilterOp
    value: Any = None
    case_sensitive: bool = False


# --------------------------------------------------------------------------- #
# Value helpers
# --------------------------------------------------------------------------- #

def to_json_value(v: Any) -> Any:
    """Convert a cell value into something JSON can carry."""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return None if math.isnan(v) or math.isinf(v) else v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, dt.datetime):
        # Excel has no separate date type; show midnight datetimes as plain dates.
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    if hasattr(v, "item"):  # numpy scalars
        return to_json_value(v.item())
    return str(v)


def is_number(v: Any) -> bool:
    return isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) and not (
        isinstance(v, float) and math.isnan(v)
    )


def _parse_number(s: Any) -> float | None:
    if is_number(s):
        return float(s)
    if isinstance(s, str):
        try:
            return float(s.strip().replace(",", ""))
        except ValueError:
            return None
    return None


def _coerce_csv(s: str) -> Any:
    if s == "":
        return None
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def _is_empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _trim(rows: Iterable[Iterable[Any]]) -> list[list[Any]]:
    """Materialize rows, dropping trailing empty cells and trailing empty rows."""
    out: list[list[Any]] = []
    for row in rows:
        r = list(row)
        while r and r[-1] is None:
            r.pop()
        out.append(r)
    while out and not out[-1]:
        out.pop()
    return out


def _as_datetime(v: dt.date) -> dt.datetime:
    return v if isinstance(v, dt.datetime) else dt.datetime.combine(v, dt.time())


def _parse_cell(cell: str) -> tuple[int, int]:
    """'B3' / '$B$3' -> (row, col), raising ExcelOpsError if it isn't a valid address."""
    try:
        letters, row = coordinate_from_string(cell.replace("$", "").upper())
        col = column_index_from_string(letters)
    except (CellCoordinatesException, ValueError, AttributeError):
        raise ExcelOpsError(f"Invalid cell reference '{cell}' (expected e.g. 'B3')") from None
    if not (1 <= row <= EXCEL_MAX_ROW and 1 <= col <= EXCEL_MAX_COL):
        raise ExcelOpsError(f"Cell '{cell}' is outside Excel's grid (max XFD1048576)")
    return row, col


def _ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


# --------------------------------------------------------------------------- #
# Operations
# --------------------------------------------------------------------------- #

class ExcelOps:
    def __init__(self, config: Config | None = None) -> None:
        self.config = config or Config()

    # ---- path & file guards ---------------------------------------------- #

    def resolve_path(self, path: str, *, must_exist: bool = True, write: bool = False) -> Path:
        if not path or "\x00" in path:
            raise ExcelOpsError("A file path is required")
        p = Path(path).expanduser()
        if not p.is_absolute():
            p = self.config.allowed_dirs[0] / p
        p = p.resolve()  # collapses '..' and follows symlinks before the sandbox check

        if not any(p.is_relative_to(root) for root in self.config.allowed_dirs):
            allowed = ", ".join(str(d) for d in self.config.allowed_dirs)
            raise ExcelOpsError(f"Access denied: {p} is outside the allowed directories ({allowed})")

        ext = p.suffix.lower()
        allowed_ext = WRITE_EXTENSIONS if write else READ_EXTENSIONS
        if ext not in allowed_ext:
            kind = "write" if write else "read"
            raise ExcelOpsError(
                f"Unsupported file type '{ext or '(none)'}' for {kind}. "
                f"Supported: {', '.join(sorted(allowed_ext))}"
            )

        if must_exist:
            if not p.is_file():
                raise ExcelOpsError(f"File not found: {p}")
            size_mb = p.stat().st_size / (1024 * 1024)
            if size_mb > self.config.max_file_mb:
                raise ExcelOpsError(
                    f"File is {size_mb:.1f} MB, over the {self.config.max_file_mb:g} MB limit"
                )
        return p

    def _require_writable(self) -> None:
        if self.config.read_only:
            raise ExcelOpsError("Server is running in read-only mode; writes are disabled")

    # ---- low-level readers ------------------------------------------------ #

    def _sheet_names(self, p: Path) -> list[str]:
        ext = p.suffix.lower()
        if ext == ".csv":
            return [p.stem]
        if ext == ".xls":
            import xlrd

            book = xlrd.open_workbook(str(p), on_demand=True)
            try:
                return book.sheet_names()
            finally:
                book.release_resources()
        wb = openpyxl.load_workbook(p, read_only=True)
        try:
            return wb.sheetnames
        finally:
            wb.close()

    def _pick_sheet(self, names: list[str], sheet: str | None) -> str:
        if sheet is None or sheet == "":
            return names[0]
        if sheet in names:
            return sheet
        for n in names:
            if n.lower() == sheet.lower():
                return n
        raise ExcelOpsError(f"Sheet '{sheet}' not found. Available sheets: {names}")

    def _read_grids(self, p: Path, sheets: list[str] | None = None) -> dict[str, list[list[Any]]]:
        """Return {sheet_name: rows} of cell values (row 1 / column A at index 0)."""
        ext = p.suffix.lower()
        if ext == ".csv":
            with open(p, newline="", encoding="utf-8-sig") as f:
                sample = f.read(64 * 1024)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                rows = ([_coerce_csv(c) for c in r] for r in csv.reader(f, dialect))
                return {p.stem: _trim(rows)}

        if ext == ".xls":
            import xlrd

            book = xlrd.open_workbook(str(p), on_demand=True)
            try:
                out = {}
                for name in sheets or book.sheet_names():
                    sh = book.sheet_by_name(name)
                    rows = []
                    for r in range(sh.nrows):
                        row = []
                        for c in range(sh.ncols):
                            cell = sh.cell(r, c)
                            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                                row.append(None)
                            elif cell.ctype == xlrd.XL_CELL_DATE:
                                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                row.append(bool(cell.value))
                            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                                row.append(int(cell.value))
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                row.append(xlrd.error_text_from_code.get(cell.value, "#ERR"))
                            else:
                                row.append(cell.value)
                        rows.append(row)
                    out[name] = _trim(rows)
                return out
            finally:
                book.release_resources()

        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            return {
                name: _trim(wb[name].iter_rows(min_row=1, min_col=1, values_only=True))
                for name in sheets or wb.sheetnames
            }
        finally:
            wb.close()

    def _read_grid(self, p: Path, sheet: str | None) -> tuple[str, list[list[Any]]]:
        name = self._pick_sheet(self._sheet_names(p), sheet)
        return name, self._read_grids(p, [name])[name]

    def _table(
        self, grid: list[list[Any]], header_row: int
    ) -> tuple[list[str], list[tuple[int, list[Any]]]]:
        """Split a grid into (column names, [(excel_row_number, values), ...])."""
        if header_row < 0:
            raise ExcelOpsError("header_row must be 0 (no header) or a 1-based row number")
        width = max((len(r) for r in grid), default=0)
        if header_row == 0:
            columns = [get_column_letter(i + 1) for i in range(width)]
        else:
            raw = grid[header_row - 1] if header_row <= len(grid) else []
            columns, seen = [], {}
            for i in range(width):
                v = raw[i] if i < len(raw) else None
                name = get_column_letter(i + 1) if _is_empty(v) else str(to_json_value(v)).strip()
                if name in seen:
                    seen[name] += 1
                    name = f"{name}_{seen[name]}"
                else:
                    seen[name] = 1
                columns.append(name)

        body = []
        for idx in range(header_row, len(grid)):
            row = grid[idx] + [None] * (width - len(grid[idx]))
            if all(_is_empty(v) for v in row):
                continue
            body.append((idx + 1, row))
        return columns, body

    def _resolve_columns(self, columns: list[str], wanted: list[str]) -> list[int]:
        idxs = []
        lower = {c.lower(): i for i, c in enumerate(columns)}
        for w in wanted:
            if w in columns:
                idxs.append(columns.index(w))
            elif w.lower() in lower:
                idxs.append(lower[w.lower()])
            else:
                raise ExcelOpsError(f"Column '{w}' not found. Available columns: {columns}")
        return idxs

    @staticmethod
    def _records(columns, body, idxs) -> list[dict[str, Any]]:
        return [
            {"_row": rownum, **{columns[i]: to_json_value(row[i]) for i in idxs}}
            for rownum, row in body
        ]

    # ---- 1. read_spreadsheet --------------------------------------------- #

    def read_spreadsheet(self, path: str, preview_rows: int = 5) -> dict[str, Any]:
        p = self.resolve_path(path)
        log.info("read_spreadsheet path=%s", p)
        preview_rows = max(0, min(preview_rows, 50))
        grids = self._read_grids(p)
        sheets = []
        for name, grid in grids.items():
            n_rows = len(grid)
            n_cols = max((len(r) for r in grid), default=0)
            sheets.append({
                "name": name,
                "rows": n_rows,
                "columns": n_cols,
                "used_range": f"A1:{_ref(n_rows, n_cols)}" if n_rows and n_cols else None,
                "preview": [[to_json_value(v) for v in r] for r in grid[:preview_rows]],
            })
        return {
            "path": str(p),
            "format": p.suffix.lower().lstrip("."),
            "size_bytes": p.stat().st_size,
            "sheet_count": len(sheets),
            "sheets": sheets,
        }

    # ---- 2. get_sheet_data ----------------------------------------------- #

    def get_sheet_data(
        self,
        path: str,
        sheet: str | None = None,
        header_row: int = 1,
        columns: list[str] | None = None,
        offset: int = 0,
        limit: int = 500,
    ) -> dict[str, Any]:
        p = self.resolve_path(path)
        log.info("get_sheet_data path=%s sheet=%s", p, sheet)
        name, grid = self._read_grid(p, sheet)
        cols, body = self._table(grid, header_row)
        idxs = self._resolve_columns(cols, columns) if columns else list(range(len(cols)))
        offset = max(0, offset)
        limit = max(1, min(limit, MAX_ROWS_PER_CALL))
        page = body[offset:offset + limit]
        return {
            "sheet": name,
            "columns": [cols[i] for i in idxs],
            "total_rows": len(body),
            "offset": offset,
            "returned": len(page),
            "truncated": offset + len(page) < len(body),
            "rows": self._records(cols, page, idxs),
        }

    # ---- 3. write_cell ---------------------------------------------------- #

    def write_cell(self, path: str, cell: str, value: Any, sheet: str | None = None) -> dict[str, Any]:
        self._require_writable()
        p = self.resolve_path(path, write=True)
        cell = _ref(*_parse_cell(cell))

        wb = self._load_for_write(p)
        ws = wb[self._pick_sheet(wb.sheetnames, sheet)]
        old = ws[cell].value
        ws[cell] = value
        self._save(wb, p)
        log.info("write_cell path=%s sheet=%s cell=%s", p, ws.title, cell)
        return {
            "path": str(p),
            "sheet": ws.title,
            "cell": cell,
            "old_value": to_json_value(old),
            "new_value": to_json_value(value),
            "is_formula": isinstance(value, str) and value.startswith("="),
        }

    # ---- 4. write_range --------------------------------------------------- #

    def write_range(
        self, path: str, start_cell: str, values: list[list[Any]], sheet: str | None = None
    ) -> dict[str, Any]:
        self._require_writable()
        p = self.resolve_path(path, write=True)
        if not values or not all(isinstance(r, list) for r in values):
            raise ExcelOpsError("values must be a non-empty 2D array (a list of rows)")
        start_row, start_col = _parse_cell(start_cell)

        wb = self._load_for_write(p)
        ws = wb[self._pick_sheet(wb.sheetnames, sheet)]
        width = max(len(r) for r in values)
        if start_row + len(values) - 1 > EXCEL_MAX_ROW or start_col + width - 1 > EXCEL_MAX_COL:
            raise ExcelOpsError("values would extend past the edge of the sheet")
        for r, row in enumerate(values):
            for c, v in enumerate(row):
                ws.cell(row=start_row + r, column=start_col + c, value=v)
        self._save(wb, p)
        end = _ref(start_row + len(values) - 1, start_col + width - 1)
        rng = f"{_ref(start_row, start_col)}:{end}"
        log.info("write_range path=%s sheet=%s range=%s", p, ws.title, rng)
        return {
            "path": str(p),
            "sheet": ws.title,
            "range": rng,
            "rows_written": len(values),
            "cells_written": sum(len(r) for r in values),
        }

    # ---- 5. create_sheet -------------------------------------------------- #

    def create_sheet(
        self,
        path: str,
        name: str,
        headers: list[str] | None = None,
        index: int | None = None,
    ) -> dict[str, Any]:
        self._require_writable()
        p = self.resolve_path(path, write=True)
        self._validate_sheet_name(name)
        wb = self._load_for_write(p)
        if name.lower() in (n.lower() for n in wb.sheetnames):
            raise ExcelOpsError(f"Sheet '{name}' already exists. Existing sheets: {wb.sheetnames}")
        ws = wb.create_sheet(title=name, index=index)
        if headers:
            self._write_headers(ws, headers)
        self._save(wb, p)
        log.info("create_sheet path=%s sheet=%s", p, name)
        return {"path": str(p), "sheet": name, "sheets": wb.sheetnames}

    # ---- 6. create_workbook ----------------------------------------------- #

    def create_workbook(
        self,
        path: str,
        sheet_name: str = "Sheet1",
        headers: list[str] | None = None,
        overwrite: bool = False,
    ) -> dict[str, Any]:
        self._require_writable()
        p = self.resolve_path(path, must_exist=False, write=True)
        if p.suffix.lower() != ".xlsx":
            raise ExcelOpsError("create_workbook only creates .xlsx files")
        if p.exists() and not overwrite:
            raise ExcelOpsError(f"File already exists: {p} (pass overwrite=true to replace it)")
        self._validate_sheet_name(sheet_name)
        p.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        if headers:
            self._write_headers(ws, headers)
        self._save(wb, p)
        log.info("create_workbook path=%s overwrite=%s", p, overwrite)
        return {"path": str(p), "sheets": [sheet_name], "headers": headers or []}

    # ---- 7. get_formulas -------------------------------------------------- #

    def get_formulas(
        self, path: str, cell_range: str | None = None, sheet: str | None = None
    ) -> dict[str, Any]:
        p = self.resolve_path(path)
        if p.suffix.lower() not in WRITE_EXTENSIONS:
            raise ExcelOpsError("Formulas can only be read from .xlsx/.xlsm files")
        log.info("get_formulas path=%s sheet=%s range=%s", p, sheet, cell_range)

        wb_f = openpyxl.load_workbook(p, read_only=True, data_only=False)
        wb_v = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            name = self._pick_sheet(wb_f.sheetnames, sheet)
            bounds: dict[str, int] = {"min_row": 1, "min_col": 1}
            if cell_range:
                try:
                    min_col, min_row, max_col, max_row = range_boundaries(cell_range.upper())
                except (CellCoordinatesException, ValueError):
                    raise ExcelOpsError(f"Invalid range '{cell_range}' (expected e.g. 'A1:D20')") from None
                bounds = {"min_row": min_row or 1, "min_col": min_col or 1}
                if max_row:
                    bounds["max_row"] = max_row
                if max_col:
                    bounds["max_col"] = max_col

            formulas = []
            rows_f = wb_f[name].iter_rows(values_only=True, **bounds)
            rows_v = wb_v[name].iter_rows(values_only=True, **bounds)
            for r, (row_f, row_v) in enumerate(zip(rows_f, rows_v), start=bounds["min_row"]):
                for c, (f, v) in enumerate(zip(row_f, row_v), start=bounds["min_col"]):
                    text = getattr(f, "text", f)  # ArrayFormula / DataTableFormula
                    if isinstance(text, str) and text.startswith("="):
                        formulas.append({
                            "cell": _ref(r, c),
                            "formula": text,
                            "cached_value": to_json_value(v),
                        })
        finally:
            wb_f.close()
            wb_v.close()
        return {"sheet": name, "range": cell_range, "formula_count": len(formulas), "formulas": formulas}

    # ---- 8. search -------------------------------------------------------- #

    def search(
        self,
        path: str,
        query: str,
        sheet: str | None = None,
        mode: Literal["contains", "exact", "regex"] = "contains",
        case_sensitive: bool = False,
        limit: int = 100,
    ) -> dict[str, Any]:
        p = self.resolve_path(path)
        log.info("search path=%s sheet=%s mode=%s", p, sheet, mode)
        names = self._sheet_names(p)
        targets = [self._pick_sheet(names, sheet)] if sheet else names
        limit = max(1, min(limit, MAX_ROWS_PER_CALL))

        flags = 0 if case_sensitive else re.IGNORECASE
        if mode == "regex":
            try:
                pattern = re.compile(query, flags)
            except re.error as e:
                raise ExcelOpsError(f"Invalid regex: {e}") from None
        q_num = _parse_number(query) if mode != "regex" else None
        q_text = query if case_sensitive else query.lower()

        def matches(v: Any) -> bool:
            if v is None:
                return False
            if q_num is not None and is_number(v) and float(v) == q_num:
                return True
            text = str(to_json_value(v))
            if mode == "regex":
                return pattern.search(text) is not None
            if not case_sensitive:
                text = text.lower()
            return text == q_text if mode == "exact" else q_text in text

        results, truncated = [], False
        for name, grid in self._read_grids(p, targets).items():
            for r, row in enumerate(grid, start=1):
                for c, v in enumerate(row, start=1):
                    if matches(v):
                        if len(results) >= limit:
                            truncated = True
                            break
                        results.append({"sheet": name, "cell": _ref(r, c), "value": to_json_value(v)})
                if truncated:
                    break
            if truncated:
                break
        return {"query": query, "mode": mode, "match_count": len(results), "truncated": truncated,
                "matches": results}

    # ---- 9. get_summary_stats --------------------------------------------- #

    def get_summary_stats(
        self,
        path: str,
        sheet: str | None = None,
        header_row: int = 1,
        columns: list[str] | None = None,
    ) -> dict[str, Any]:
        p = self.resolve_path(path)
        log.info("get_summary_stats path=%s sheet=%s", p, sheet)
        name, grid = self._read_grid(p, sheet)
        cols, body = self._table(grid, header_row)
        idxs = self._resolve_columns(cols, columns) if columns else list(range(len(cols)))

        numeric, other = {}, {}
        for i in idxs:
            vals = [row[i] for _, row in body]
            present = [v for v in vals if not _is_empty(v)]
            nums = [float(v) for v in present if is_number(v)]
            missing = len(vals) - len(present)
            # Treat a column as numeric when most of its non-empty cells are numbers.
            if nums and len(nums) * 2 >= len(present):
                numeric[cols[i]] = {
                    "count": len(nums),
                    "missing": missing,
                    "non_numeric": len(present) - len(nums),
                    "min": min(nums),
                    "max": max(nums),
                    "mean": statistics.fmean(nums),
                    "median": statistics.median(nums),
                    "stdev": statistics.stdev(nums) if len(nums) > 1 else None,
                    "sum": math.fsum(nums),
                }
            else:
                if present and all(isinstance(v, (dt.date, dt.datetime)) for v in present):
                    other[cols[i]] = {
                        "count": len(present),
                        "missing": missing,
                        "earliest": to_json_value(min(present, key=_as_datetime)),
                        "latest": to_json_value(max(present, key=_as_datetime)),
                    }
                    continue
                texts = [str(to_json_value(v)) for v in present]
                top = max(set(texts), key=texts.count) if texts else None
                other[cols[i]] = {
                    "count": len(present),
                    "missing": missing,
                    "unique": len(set(texts)),
                    "top": top,
                    "top_count": texts.count(top) if top is not None else 0,
                }
        return {"sheet": name, "row_count": len(body), "numeric_columns": numeric,
                "non_numeric_columns": other}

    # ---- 10. apply_filter ------------------------------------------------- #

    def apply_filter(
        self,
        path: str,
        conditions: list[Condition],
        sheet: str | None = None,
        header_row: int = 1,
        match: Literal["all", "any"] = "all",
        columns: list[str] | None = None,
        limit: int = 500,
    ) -> dict[str, Any]:
        p = self.resolve_path(path)
        if not conditions:
            raise ExcelOpsError("At least one condition is required")
        log.info("apply_filter path=%s sheet=%s conditions=%d", p, sheet, len(conditions))
        name, grid = self._read_grid(p, sheet)
        cols, body = self._table(grid, header_row)
        cond_idx = [self._resolve_columns(cols, [c.column])[0] for c in conditions]
        out_idx = self._resolve_columns(cols, columns) if columns else list(range(len(cols)))
        for c in conditions:
            if c.op == "regex":
                try:
                    re.compile(str(c.value))
                except re.error as e:
                    raise ExcelOpsError(f"Invalid regex for column '{c.column}': {e}") from None
            if c.op == "between" and not (isinstance(c.value, list) and len(c.value) == 2):
                raise ExcelOpsError("'between' needs value=[low, high]")
            if c.op in ("in", "not_in") and not isinstance(c.value, list):
                raise ExcelOpsError(f"'{c.op}' needs value to be a list")

        combine = all if match == "all" else any
        hits = [
            (rownum, row) for rownum, row in body
            if combine(_eval(c, row[i]) for c, i in zip(conditions, cond_idx))
        ]
        limit = max(1, min(limit, MAX_ROWS_PER_CALL))
        return {
            "sheet": name,
            "columns": [cols[i] for i in out_idx],
            "total_rows": len(body),
            "match_count": len(hits),
            "returned": min(len(hits), limit),
            "truncated": len(hits) > limit,
            "rows": self._records(cols, hits[:limit], out_idx),
        }

    # ---- write helpers ---------------------------------------------------- #

    @staticmethod
    def _validate_sheet_name(name: str) -> None:
        if not name or len(name) > 31 or INVALID_SHEET_CHARS.search(name):
            raise ExcelOpsError(
                f"Invalid sheet name '{name}': must be 1-31 characters without [ ] : * ? / \\"
            )

    @staticmethod
    def _write_headers(ws, headers: list[str]) -> None:
        bold = Font(bold=True)
        for c, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c, value=h).font = bold
        ws.freeze_panes = "A2"

    @staticmethod
    def _load_for_write(p: Path):
        try:
            return openpyxl.load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
        except PermissionError:
            raise ExcelOpsError(f"Cannot open {p}: it may be open in another program") from None

    @staticmethod
    def _save(wb, p: Path) -> None:
        """Save atomically: write a temp file beside the target, then swap it in."""
        fd, tmp = tempfile.mkstemp(dir=p.parent, prefix=f".{p.stem}-", suffix=p.suffix)
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, p)
        except PermissionError:
            raise ExcelOpsError(f"Cannot save {p}: it may be open in another program") from None
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)


# --------------------------------------------------------------------------- #
# Filter evaluation
# --------------------------------------------------------------------------- #

def _comparable(cell: Any, target: Any) -> tuple[Any, Any] | None:
    """Coerce a cell and a filter value to a common orderable type, if possible."""
    if isinstance(cell, (dt.datetime, dt.date)) and isinstance(target, str):
        try:
            t = dt.datetime.fromisoformat(target)
        except ValueError:
            return None
        return _as_datetime(cell), t
    cn, tn = _parse_number(cell), _parse_number(target)
    if cn is not None and tn is not None:
        return cn, tn
    if isinstance(cell, str) and isinstance(target, str):
        return cell.lower(), target.lower()
    return None


def _text(v: Any, case_sensitive: bool) -> str:
    s = "" if v is None else str(to_json_value(v))
    return s if case_sensitive else s.lower()


def _equal(cell: Any, target: Any, case_sensitive: bool) -> bool:
    if cell is None or target is None:
        return _is_empty(cell) and _is_empty(target)
    if isinstance(cell, bool) or isinstance(target, bool):
        return cell == target
    pair = _comparable(cell, target)
    if pair is not None and not isinstance(pair[0], str):
        return pair[0] == pair[1]
    return _text(cell, case_sensitive) == _text(target, case_sensitive)


def _order(cell: Any, target: Any, op: str) -> bool:
    pair = _comparable(cell, target)
    if pair is None:
        return False
    a, b = pair
    return {"gt": a > b, "gte": a >= b, "lt": a < b, "lte": a <= b}[op]


def _eval(c: Condition, cell: Any) -> bool:
    op, val, cs = c.op, c.value, c.case_sensitive
    if op == "is_null":
        return _is_empty(cell)
    if op == "not_null":
        return not _is_empty(cell)
    if op == "eq":
        return _equal(cell, val, cs)
    if op == "ne":
        return not _equal(cell, val, cs)
    if op in ("gt", "gte", "lt", "lte"):
        return cell is not None and _order(cell, val, op)
    if op == "between":
        return cell is not None and _order(cell, val[0], "gte") and _order(cell, val[1], "lte")
    if op == "in":
        return any(_equal(cell, v, cs) for v in val)
    if op == "not_in":
        return not any(_equal(cell, v, cs) for v in val)
    if cell is None:
        return op == "not_contains"
    text, needle = _text(cell, cs), _text(val, cs)
    if op == "contains":
        return needle in text
    if op == "not_contains":
        return needle not in text
    if op == "startswith":
        return text.startswith(needle)
    if op == "endswith":
        return text.endswith(needle)
    if op == "regex":
        return re.search(str(val), str(to_json_value(cell)), 0 if cs else re.IGNORECASE) is not None
    raise ExcelOpsError(f"Unknown operator '{op}'")
